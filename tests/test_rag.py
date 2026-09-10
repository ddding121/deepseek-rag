import io
import unittest
import numpy as np
from docx import Document
from pypdf import PdfWriter
from rag import Chunk, parse_file, split_text, retrieve, build_messages, answer


class FakeEncoder:
    def encode(self, texts, **kwargs):
        return np.array([[1., 0.]])


class RagTests(unittest.TestCase):
    def test_docx_paragraph_and_table(self):
        doc = Document()
        doc.add_paragraph('试用期请假须直属经理审批。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '年假五天'
        data = io.BytesIO()
        doc.save(data)
        chunks, _ = parse_file('制度.docx', data.getvalue())
        self.assertEqual(len(chunks), 2)
        self.assertIn('表格', chunks[1].location)

    def test_empty_pdf_and_legacy_word(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaises(ValueError):
            parse_file('空白.pdf', data.getvalue())
        with self.assertRaises(ValueError):
            parse_file('旧版.doc', b'bad')

    def test_overlap_no_missing_characters(self):
        parts = list(split_text('abcdefghijk', 5, 2))
        self.assertEqual(parts, ['abcde', 'defgh', 'ghijk'])

    def test_ranking_and_threshold(self):
        chunks = [Chunk('a', '1', 'a'), Chunk('b', '2', 'b')]
        hits = retrieve('问题', chunks, np.array([[0., 1.], [1., 0.]]), FakeEncoder())
        self.assertEqual(hits[0][0].source, 'b')
        self.assertEqual(len(hits), 1)
        self.assertEqual(answer('问题', [], '', ''), '知识库中没有找到足够依据。')
        self.assertIn('"id": 1', build_messages('问题', hits)[1]['content'])


if __name__ == '__main__':
    unittest.main()
